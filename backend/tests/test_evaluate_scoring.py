"""F15 scoring harness: labels of spec 10.2, metrics, overrides, error log, workbook, CLI."""
import csv
import json

import pytest

from analyser.evaluate import (export_scores_xlsx, load_golden, score_run, score_test,
                               write_error_log)
from analyser.evaluate.__main__ import main as cli_main
from analyser.evaluate.golden_check import golden_answer, golden_envelopes
from analyser.models import AnswerEnvelope, ReviewState, Verification

GOLDEN = {g["test_id"]: g for g in load_golden()}


def tool(test_id: str, **changes) -> AnswerEnvelope:
    """The expected answer of a golden test, with some fields changed (a fake tool answer)."""
    a = golden_answer(GOLDEN[test_id]).model_copy(update=changes)
    return AnswerEnvelope(question="q", answer=a)


def label(test_id: str, env: AnswerEnvelope | None):
    return score_test(GOLDEN[test_id], env)


# --------------------------------------------------------------------------- sanity


def test_golden_scored_against_itself_is_100_percent():
    s = score_run(golden_envelopes(), "sanity", overrides=[])
    assert len(s.tests) == 208
    assert s.counts["Correct"] + s.counts["Correct absent"] == 208
    assert s.counts["Correct absent"] == 16
    assert s.precision == 1.0 and s.recall == 1.0 and s.core_pass_rate == 1.0
    assert s.hallucinations == 0
    assert set(s.by_contract) == {"BM1", "BM2", "BM3", "BM4"} and set(s.by_clause) == {f"C{i}" for i in range(7)}


# --------------------------------------------------------------------------- labels


def test_wrong_number_is_incorrect():
    t = label("BM1-2.2", tool("BM1-2.2", answer_value="8% of every payment for Works"))
    assert t.label == "Incorrect" and "wrong number" in t.reason
    assert t.cause_code == "RATE_CAP_MIXED"


def test_missing_one_of_three_numbers_is_partial():
    # expected: IPC within 15 days; within 10 days a 90% part payment
    t = label("BM1-1.2", tool("BM1-1.2", answer_value="Interim Payment Certificate within 15 days; "
                                                     "90% part payment recommended"))
    assert t.label == "Partial" and "10 days" in t.reason
    assert t.cause_code == "ONE_OF_MANY_NUMBERS"


def test_not_stated_where_found_expected_is_missed():
    t = label("BM1-2.2", tool("BM1-2.2", status="NOT_STATED", answer_value="", verbatim_quote=""))
    assert t.label == "Missed" and t.cause_code == "RETRIEVAL_MISSED"
    t2 = label("BM2-3.1", tool("BM2-3.1", status="NOT_STATED", answer_value="", verbatim_quote=""))
    assert t2.label == "Missed" and t2.cause_code == "TEXT_LAYER_MISSED_STRIKE"
    t3 = label("BM1-2.2", None)                       # no answer at all for the question
    assert t3.label == "Missed"


def test_value_where_not_stated_expected_is_wrongly_present():
    env = tool("BM1-1.8", status="FOUND", answer_value="Pay when paid: 30 days after funds arrive",
               verbatim_quote="x", pdf_page=10)
    t = label("BM1-1.8", env)
    assert t.label == "Wrongly present" and t.cause_code == "INVENTED_VALUE"
    s = score_run({"BM1": [env]}, "r", overrides=[])
    assert s.hallucinations == 1


def test_both_not_stated_is_correct_absent():
    assert label("BM1-1.8", tool("BM1-1.8")).label == "Correct absent"


def test_accept_also_status_is_partial():
    t = label("BM4-1.3", tool("BM4-1.3", status="NOT_STATED", answer_value="", verbatim_quote=""))
    assert t.label == "Partial" and "accept_also" in t.reason


def test_wrong_status_is_incorrect():
    t = label("BM1-4.3", tool("BM1-4.3", status="POINTS_ELSEWHERE"))
    assert t.label == "Incorrect" and t.cause_code == "REFERENCE_NOT_FOLLOWED"


def test_formatting_differences_are_ignored():
    t1 = label("BM2-6.2", tool("BM2-6.2", answer_value="PG 5 % of tendered amount; SD 5 % of tendered "
                                                       "value; EMD Rs 424547; BG 110 % of the advance"))
    t2 = label("BM1-3.2", tool("BM1-3.2", answer_value="10 % of Contract Price, 2 instalments of 5 %"))
    t3 = label("BM2-1.6", tool("BM2-1.6", answer_value="Gross work of Rs 3,000,000 needed for an interim bill"))
    assert (t1.label, t2.label, t3.label) == ("Correct", "Correct", "Correct")


def test_rate_conversion_only_counts_when_stated():
    stated = tool("BM2-4.2", answer_value="0.0333% of accepted tendered value per day "
                                          "(= 1% per month / 30), set by the Schedule F authority")
    assert label("BM2-4.2", stated).label == "Correct"
    unstated = tool("BM2-4.2", answer_value="0.0333% of accepted tendered value per day")
    assert label("BM2-4.2", unstated).label == "Incorrect"


def test_wrong_base_and_missing_base():
    t = label("BM1-2.3", tool("BM1-2.3", answer_value="5% of Bid Price", unit="% of Bid Price"))
    assert t.label == "Incorrect" and "base" in t.reason
    t2 = label("BM1-2.3", tool("BM1-2.3", answer_value="5% cap", unit="%"))
    assert t2.label == "Partial" and t2.cause_code == "WRONG_BASE_OR_UNIT"


def test_yes_no_polarity():
    t = label("BM1-2.1", tool("BM1-2.1", answer_value="No retention is deducted"))
    assert t.label == "Incorrect" and t.cause_code == "NEGATIVE_MISREAD"


def test_text_answers_are_flagged_for_a_person():
    t = label("BM1-5.5", tool("BM1-5.5", answer_value="Base date is the last day of the month at least "
                                                     "28 days before the bid due date"))
    assert t.label == "Correct" and t.needs_human_check


def test_quote_not_verified_is_incorrect():
    env = tool("BM1-2.2")
    env.verification = Verification(question_id="2.2", quote_found=False, verdict="fail", discard=True)
    t = label("BM1-2.2", env)
    assert t.label == "Incorrect" and t.quote_verified is False and t.cause_code == "INVENTED_QUOTE"


def test_review_edit_counts():
    env = tool("BM1-2.2", answer_value="8% of every payment for Works")
    fixed = env.answer.model_copy(update={"answer_value": "6% of every payment for Works"})
    env.review = ReviewState(decision="edited", user_answer=fixed)
    assert label("BM1-2.2", env).label == "Correct"


# --------------------------------------------------------------------------- run level


def test_metrics_on_a_mixed_run():
    envs = golden_envelopes()["BM1"]
    by_q = {e.answer.question_id: e for e in envs}
    by_q["2.2"].answer.answer_value = "8% of every payment"                  # Incorrect
    by_q["4.2"].answer.status = "NOT_STATED"                                  # Missed
    s = score_run({"BM1": envs}, "mixed", overrides=[])
    assert s.counts["Incorrect"] == 1 and s.counts["Missed"] == 1
    assert s.counts["total"] == 52
    given = 52 - 5 - 1                    # 5 expected NOT_STATED + the missed one
    assert s.precision == round(45 / given, 4)
    assert s.recall == round(45 / 47, 4)
    core = [t for t in s.tests if t.core]
    assert s.core_pass_rate == round((len(core) - 2) / len(core), 4)


def test_manual_override_applied_last_and_marked():
    env = tool("BM1-2.2", answer_value="8% of every payment for Works")
    overrides = [{"test_id": "BM1-2.2", "label": "Correct", "reason": "checked by hand", "by": "Asha"},
                 {"test_id": "BM1-2.2", "label": "Missed", "reason": "other run", "run_id": "old"},
                 {"test_id": "BM1-2.2", "label": "Not a label", "reason": "ignored"}]
    s = score_run({"BM1": [env]}, "r1", overrides=overrides)
    t = next(t for t in s.tests if t.test_id == "BM1-2.2")
    assert t.label == "Correct" and t.reason.startswith("MANUAL OVERRIDE by Asha")
    assert "auto label Incorrect" in t.reason


def test_overrides_file_is_valid():
    from analyser.evaluate.scoring import OVERRIDES_FILE, load_overrides
    assert OVERRIDES_FILE.exists() and isinstance(load_overrides(), list)


def test_error_log_csv(tmp_path):
    envs = golden_envelopes()["BM1"]
    envs[15].answer.answer_value = "8% of every payment"          # BM1-2.2
    s = score_run({"BM1": envs}, "run-x", overrides=[])
    path = tmp_path / "errors.csv"
    write_error_log(s, path)
    rows = list(csv.DictReader(path.open(encoding="utf-8-sig")))
    assert list(rows[0])[:9] == ["run_id", "contract", "question", "expected", "tool_output", "label",
                                 "cause_code", "fix_tried", "result"]
    assert [r["test_id"] for r in rows] == ["BM1-2.2"] and rows[0]["label"] == "Incorrect"
    # a person's notes survive a rewrite
    rows[0]["fix_tried"] = "prompt v2"
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0]))
        w.writeheader()
        w.writerows(rows)
    write_error_log(s, path)
    assert next(csv.DictReader(path.open(encoding="utf-8-sig")))["fix_tried"] == "prompt v2"


def test_xlsx_export(tmp_path):
    from openpyxl import load_workbook
    s = score_run(golden_envelopes(), "sanity", overrides=[])
    path = tmp_path / "scores.xlsx"
    export_scores_xlsx(s, path)
    wb = load_workbook(path)
    assert wb.sheetnames == ["Summary", "Tests", "Error log"]
    assert wb["Tests"].max_row == 209


def test_cli(tmp_path, capsys):
    run = tmp_path / "run1"
    run.mkdir()
    envs = golden_envelopes()["BM2"]
    (run / "answers.json").write_text(json.dumps([e.model_dump() for e in envs]), encoding="utf-8")
    (run / "meta.json").write_text(json.dumps({"run_id": "run1", "contract_id": "BM2"}), encoding="utf-8")
    bare = tmp_path / "bm4_answers.json"           # bare Answer records also work
    bare.write_text(json.dumps([e.answer.model_dump() for e in golden_envelopes()["BM4"]]), encoding="utf-8")
    out = tmp_path / "score.json"
    assert cli_main([str(run), str(bare), "--save", str(out), "--no-overrides"]) == 0
    text = capsys.readouterr().out
    assert "BM2" in text and "BM4" in text and "Precision" in text
    assert json.loads(out.read_text(encoding="utf-8"))["counts"]["total"] == 104


@pytest.mark.parametrize("test_id", ["BM3-1.2", "BM2-6.8"])
def test_not_stated_with_empty_expected_answer(test_id):
    assert label(test_id, tool(test_id)).label == "Correct absent"
