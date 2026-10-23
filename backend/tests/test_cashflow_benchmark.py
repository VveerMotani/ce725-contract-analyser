"""§8.4 benchmark workbook: real Excel formulas, and the engine matches the spec table typed into it."""
import openpyxl
import pytest

from analyser.cashflow.benchmark import SPEC_TABLE, WORKBOOK_PATH, engine_figures, write_workbook


@pytest.fixture(scope="module")
def wb(tmp_path_factory):
    path = write_workbook(tmp_path_factory.mktemp("bm") / "benchmark_8_4.xlsx")
    return openpyxl.load_workbook(path)          # formulas, not cached values


def test_sheets_and_manual_check_note(wb):
    assert wb.sheetnames == ["Read me", "Inputs", "Benchmark 8.4", "Monthly", "Python monthly"]
    text = " ".join(str(c.value) for row in wb["Read me"].iter_rows() for c in row if c.value)
    assert "checked by a person" in text
    assert "Checked by:" in text


def test_monthly_sheet_uses_formulas(wb):
    ws = wb["Monthly"]
    for row in ws.iter_rows(min_row=2, max_row=32, min_col=2, max_col=20):
        for c in row:
            assert isinstance(c.value, str) and c.value.startswith("="), c.coordinate
    bm = wb["Benchmark 8.4"]
    for r in range(2, 2 + len(SPEC_TABLE)):
        assert str(bm.cell(row=r, column=4).value).startswith("=")


def test_engine_matches_spec_table_in_workbook(wb):
    """Column C holds the spec §8.4 values typed by hand; the engine must match each to 4 decimals."""
    bm = wb["Benchmark 8.4"]
    py = engine_figures()
    for r, (key, item, hand, val) in enumerate(SPEC_TABLE, start=2):
        typed = bm.cell(row=r, column=3).value
        assert typed == val, item
        assert round(py[key], 4) == pytest.approx(typed, abs=1e-12), item


def test_spec_table_is_the_printed_one():
    vals = {k: v for k, _, _, v in SPEC_TABLE}
    assert vals == {"gross_month": 4.1667, "retention_month": 0.25, "retention_cap_month": 20,
                    "recovery_month": 0.625, "advance_repaid_month": 16, "deadline_month": 19.2,
                    "net_month1": 3.2917, "ld_60": 3.0, "ld_days_to_cap": 200, "ld_at_cap": 10.0}


def test_committed_workbook_exists():
    assert WORKBOOK_PATH.exists(), "run: python -m analyser.cashflow.benchmark"
