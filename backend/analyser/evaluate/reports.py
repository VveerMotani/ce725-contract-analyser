"""Outputs of the evaluation: error log CSV (spec 10.6), Excel workbook, text table."""
from __future__ import annotations

import csv
from pathlib import Path

from analyser.models import ScoreSummary

from .scoring import LABELS

ERROR_COLUMNS = ["run_id", "contract", "question", "expected", "tool_output", "label",
                 "cause_code", "fix_tried", "result"]
EXTRA_COLUMNS = ["test_id", "reason", "needs_human_check"]
ERROR_LABELS = {"Incorrect", "Missed", "Wrongly present", "Partial"}


def _old_entries(path: Path) -> dict[tuple[str, str], dict]:
    """Rows a person already filled in (fix tried, result, cause code) are kept."""
    if not path.exists():
        return {}
    with path.open(newline="", encoding="utf-8-sig") as f:
        return {(r.get("run_id", ""), r.get("test_id", "")): r for r in csv.DictReader(f)}


def write_error_log(summary: ScoreSummary, path: Path, include_partial: bool = True) -> None:
    """One row per Incorrect, Missed and Wrongly present test (and Partial, reported separately).

    cause_code is the heuristic suggestion unless a person has edited it in an earlier file.
    """
    path = Path(path)
    old = _old_entries(path)
    labels = ERROR_LABELS if include_partial else ERROR_LABELS - {"Partial"}
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=ERROR_COLUMNS + EXTRA_COLUMNS)
        w.writeheader()
        for t in summary.tests:
            if t.label not in labels:
                continue
            prev = old.get((summary.run_id, t.test_id), {})
            w.writerow({
                "run_id": summary.run_id,
                "contract": t.contract_id,
                "question": t.question_id,
                "expected": f"{t.expected_status}: {t.expected_answer}",
                "tool_output": f"{t.tool_status}: {t.tool_answer}",
                "label": t.label,
                "cause_code": prev.get("cause_code") or t.cause_code,
                "fix_tried": prev.get("fix_tried", ""),
                "result": prev.get("result", ""),
                "test_id": t.test_id,
                "reason": t.reason,
                "needs_human_check": "yes" if t.needs_human_check else "",
            })


def metrics_rows(summary: ScoreSummary) -> list[tuple[str, object]]:
    c = summary.counts
    judged = sum(t.reason.startswith("LLM-judged:") for t in summary.tests)
    overridden = sum(t.reason.startswith("MANUAL OVERRIDE") for t in summary.tests)
    return [
        ("Run", summary.run_id),
        ("Tests scored", c.get("total", 0)),
        ("Precision (Correct / answers given)", summary.precision),
        ("Recall (Correct / answers in key)", summary.recall),
        ("Core pass rate", summary.core_pass_rate),
        ("Quote verification rate", summary.quote_verification_rate),
        ("Hallucinations (Wrongly present + invented numbers)", summary.hallucinations),
        ("Headline correct (Correct + Correct absent) / total",
         round((c.get("Correct", 0) + c.get("Correct absent", 0)) / c["total"], 4) if c.get("total") else 0.0),
        ("Needs human check", sum(t.needs_human_check for t in summary.tests)),
        ("Labels set by code (deterministic)", c.get("total", 0) - judged - overridden),
        ("Labels set by LLM judge", judged),
        ("Labels set by manual override", overridden),
    ]


def format_table(summary: ScoreSummary) -> str:
    """Plain-text metrics table for the terminal."""
    cols = LABELS + ["total"]
    short = {"Correct absent": "CorrAbs", "Wrongly present": "WrongPres", "Incorrect": "Incorr"}
    head = f"{'':10}" + "".join(f"{short.get(c, c):>10}" for c in cols)
    lines = [f"Run: {summary.run_id}", "", "By contract", head]
    for name, counts in list(summary.by_contract.items()) + [("ALL", summary.counts)]:
        lines.append(f"{name:10}" + "".join(f"{counts.get(c, 0):>10}" for c in cols))
    lines += ["", "By clause type", head]
    for name, counts in summary.by_clause.items():
        lines.append(f"{name:10}" + "".join(f"{counts.get(c, 0):>10}" for c in cols))
    lines.append("")
    for k, v in metrics_rows(summary)[1:]:
        lines.append(f"{k:55} {v:.1%}" if isinstance(v, float) else f"{k:55} {v}")
    return "\n".join(lines)


def export_scores_xlsx(summary: ScoreSummary, path: Path) -> None:
    """Workbook for the validation report: Summary, Tests and Error log sheets."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill

    bold = Font(bold=True)
    head_fill = PatternFill("solid", fgColor="DDEBF7")
    wb = Workbook()
    ws = wb.active
    ws.title = "Summary"
    for k, v in metrics_rows(summary):
        ws.append([k, v])
        if isinstance(v, float):
            ws.cell(ws.max_row, 2).number_format = "0.0%"
    for title, table in (("By contract", summary.by_contract), ("By clause type", summary.by_clause)):
        ws.append([])
        ws.append([title])
        ws.cell(ws.max_row, 1).font = bold
        ws.append([""] + LABELS + ["total"])
        for cell in ws[ws.max_row]:
            cell.font, cell.fill = bold, head_fill
        for name, counts in list(table.items()) + [("ALL", summary.counts)]:
            ws.append([name] + [counts.get(c, 0) for c in LABELS + ["total"]])
    ws.column_dimensions["A"].width = 52
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 14

    fields = ["test_id", "contract_id", "question_id", "clause_id", "core", "expected_status",
              "tool_status", "label", "expected_answer", "tool_answer", "reason", "cause_code",
              "quote_verified", "needs_human_check"]
    wt = wb.create_sheet("Tests")
    wt.append(fields)
    for t in summary.tests:
        d = t.model_dump()
        wt.append(["" if d[f] is None else d[f] for f in fields])
    _style_sheet(wt, bold, head_fill, {"I": 50, "J": 50, "K": 50})

    we = wb.create_sheet("Error log")
    we.append(ERROR_COLUMNS + EXTRA_COLUMNS)
    for t in summary.tests:
        if t.label in ERROR_LABELS:
            we.append([summary.run_id, t.contract_id, t.question_id,
                       f"{t.expected_status}: {t.expected_answer}", f"{t.tool_status}: {t.tool_answer}",
                       t.label, t.cause_code, "", "", t.test_id, t.reason,
                       "yes" if t.needs_human_check else ""])
    _style_sheet(we, bold, head_fill, {"D": 50, "E": 50, "K": 50})
    Path(path).parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)


def _style_sheet(ws, bold, fill, widths: dict[str, int]) -> None:
    for cell in ws[1]:
        cell.font, cell.fill = bold, fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
