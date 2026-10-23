"""Spec §8.4 hand benchmark: inputs, the spec's table, and an Excel workbook built
from live Excel formulas (not values) so the team can check the arithmetic in
Excel independently of Python.

Run from backend/:  python -m analyser.cashflow.benchmark
Writes:             analyser/cashflow/benchmark_8_4.xlsx
"""
from __future__ import annotations

from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.workbook.defined_name import DefinedName

from ..models import AdvanceInstalment, CashInputs, ReleaseStep
from .engine import ld_amount, ld_days_to_cap, milestones, run_cashflow

CRORE = 1e7
WORKBOOK_PATH = Path(__file__).resolve().parent / "benchmark_8_4.xlsx"

# The §8.4 table exactly as printed in the spec (Rs crore unless stated).
SPEC_TABLE: list[tuple[str, str, str, float]] = [
    # key, item, hand calculation, result
    ("gross_month", "Gross work each month (Rs crore)", "100 / 24", 4.1667),
    ("retention_month", "Retention each month (Rs crore)", "6% of 4.1667", 0.25),
    ("retention_cap_month", "Month in which retention reaches its 5% cap (Rs 5 crore)", "5 / 0.25", 20),
    ("recovery_month", "Advance recovery each month (Rs crore)", "15% of 4.1667", 0.625),
    ("advance_repaid_month", "Month in which a Rs 10 crore advance is fully recovered", "10 / 0.625", 16),
    ("deadline_month", "80% recovery limit (month)", "0.8 x 24", 19.2),
    ("net_month1", "Net certified in month 1 (Rs crore)", "4.1667 - 0.25 - 0.625", 3.2917),
    ("ld_60", "Delay damages for 60 days (Rs crore)", "60 x 0.05% of 100", 3.0),
    ("ld_days_to_cap", "Days of delay to reach the cap", "10% / 0.05%", 200),
    ("ld_at_cap", "Delay damages at the cap (Rs crore)", "10% of 100", 10.0),
]


def benchmark_inputs() -> CashInputs:
    """The §8.4 case. Items the spec does not fix are [ASSUMPTION]s chosen to keep months whole:
    advance instalments in months 0 and 2, payment 30 days, retention released 1 month after
    completion, delay 60 days, cost share 85%, overdraft 10% a year."""
    return CashInputs(
        contract_value=100 * CRORE, duration_months=24, progress=None,
        cost_share=0.85, overdraft_rate_pa=0.10, delay_days=60,
        cert_lag_days=0, pay_lag_days=30,
        retention_rate=0.06, retention_cap=0.05,
        retention_release=[ReleaseStep(share=1.0, months_after_completion=1.0)],
        advance_instalments=[AdvanceInstalment(share_of_value=0.05, month=0),
                             AdvanceInstalment(share_of_value=0.05, month=2)],
        advance_recovery_rate=0.15, advance_recovery_deadline=0.8,
        ld_rate_per_day=0.0005, ld_cap=0.10,
    )


def engine_figures(inp: CashInputs | None = None) -> dict[str, float]:
    """The engine's numbers for each SPEC_TABLE key, in the table's units (Rs crore or months/days)."""
    inp = inp or benchmark_inputs()
    res = run_cashflow(inp, sensitivities=False)
    m1 = res.months[1]
    ms = milestones(inp, res.months)
    return {
        "gross_month": m1.gross / CRORE,
        "retention_month": m1.retention / CRORE,
        "retention_cap_month": ms["retention_cap_month"],
        "recovery_month": m1.advance_recovery / CRORE,
        "advance_repaid_month": ms["advance_repaid_month"],
        "deadline_month": inp.advance_recovery_deadline * inp.duration_months,
        "net_month1": m1.net_certified / CRORE,
        "ld_60": ld_amount(inp, 60) / CRORE,
        "ld_days_to_cap": ld_days_to_cap(inp),
        "ld_at_cap": ld_amount(inp, 10_000) / CRORE,
        # beyond the spec table
        "peak_funding": res.peak_funding / CRORE,
        "peak_month": res.peak_month,
        "overdraft_interest": res.overdraft_interest / CRORE,
        "total_cash_in": res.months[-1].cum_in / CRORE,
        "total_cash_out": res.months[-1].cum_out / CRORE,
    }


# ---------------------------------------------------------------------------
# Workbook
# ---------------------------------------------------------------------------

BOLD = Font(bold=True)
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")
INPUT_FILL = PatternFill("solid", fgColor="FFF2CC")
WARN_FILL = PatternFill("solid", fgColor="F8CBAD")

INPUTS = [
    # name, label, value, unit
    ("ContractValue", "Contract value V", 100, "Rs crore"),
    ("Months", "Duration T", 24, "months (even progress, 1/T each month)"),
    ("RetRate", "Retention rate r", 0.06, "share of each gross bill"),
    ("RetCap", "Retention cap Rmax", 0.05, "share of V"),
    ("AdvFirst", "Advance instalment 1", 0.05, "share of V"),
    ("AdvFirstMonth", "Instalment 1 arrives in month", 0, "month [ASSUMPTION]"),
    ("AdvSecond", "Advance instalment 2", 0.05, "share of V"),
    ("AdvSecondMonth", "Instalment 2 arrives in month", 2, "month [ASSUMPTION]"),
    ("RecRate", "Advance recovery rate rho", 0.15, "share of each gross bill"),
    ("RecDeadline", "Recovery deadline", 0.8, "share of T"),
    ("LDRate", "Delay damages rate l", 0.0005, "share of V per day"),
    ("LDCap", "Delay damages cap Lmax", 0.10, "share of V"),
    ("DelayDays", "Delay D", 60, "days"),
    ("CostShare", "Cost share c", 0.85, "share of gross work [ASSUMPTION]"),
    ("PayLagDays", "Payment lag", 30, "days; whole months only in this sheet (30, 60, ...)"),
    ("RelMonths", "Retention released after completion", 1, "months; whole months only [ASSUMPTION]"),
    ("ODRate", "Overdraft rate", 0.10, "per year [ASSUMPTION]"),
]

MONTH_COLS = [
    # header, formula template ({r} = this row, {p} = previous row)
    ("Month", None),
    ("Gross work B", "=IF(AND(A{r}>=1,A{r}<=Months),ContractValue/Months,0)"),
    ("Retention Ret", "=MAX(0,MIN(RetRate*B{r},RetCap*ContractValue-D{p}))"),
    ("Cum retention", "=D{p}+C{r}"),
    ("Advance received", "=IF(A{r}=AdvFirstMonth,AdvFirst*ContractValue,0)+IF(A{r}=AdvSecondMonth,AdvSecond*ContractValue,0)"),
    ("Cum advance received", "=F{p}+E{r}"),
    ("Advance recovery Rec",
     "=IF(A{r}=0,0,IF(OR(A{r}>=Months,A{r}>=MAX(1,FLOOR(RecDeadline*Months,1))),F{r}-H{p},"
     "MIN(RecRate*B{r},F{r}-H{p})))"),
    ("Cum recovered", "=H{p}+G{r}"),
    ("Advance outstanding", "=F{r}-H{r}"),
    ("Net certified", "=B{r}-C{r}-G{r}"),
    ("Retention released", "=IF(A{r}=Months+DelayDays/30+RelMonths,MAX($D$2:$D$40),0)"),
    ("Delay damages deducted", "=IF(A{r}=Months+DelayDays/30,'Benchmark 8.4'!$D${ld},0)"),
    ("Cash in", "=IF(A{r}-PayLagDays/30>=0,OFFSET(J{r},-PayLagDays/30,0),0)+E{r}+K{r}-L{r}"),
    ("Cash out", "=CostShare*B{r}"),
    ("Cum in", "=O{p}+M{r}"),
    ("Cum out", "=P{p}+N{r}"),
    ("Gap (out - in)", "=P{r}-O{r}"),
    ("Retention held", "=D{r}-SUM($K$2:K{r})"),
    ("Cap reached? (month or 999)", "=IF(D{r}>=RetCap*ContractValue-0.000001,A{r},999)"),
    ("Advance repaid? (month or 999)", "=IF(AND(A{r}>=1,H{r}>=(AdvFirst+AdvSecond)*ContractValue-0.000001),A{r},999)"),
]
FIRST_ROW, LAST_MONTH = 2, 30
LD_ROW = 2 + [k for k, *_ in SPEC_TABLE].index("ld_60")   # row of 'Benchmark 8.4' holding damages for D days


def _readme(ws) -> None:
    lines = [
        ("Spec §8.4 hand benchmark (CE725 Team 2, module M4 cash flow)", True),
        ("", False),
        ("MANUAL CHECK REQUIRED. The course requires this hand benchmark to be checked by a person.", True),
        ("Open the 'Benchmark 8.4' sheet, redo each line by hand or calculator, and sign below.", False),
        ("Checked by: ____________________   Date: ____________   Result (all MATCH / differences): ______________", False),
        ("", False),
        ("What this file is", True),
        ("Every number in 'Monthly' and in column D of 'Benchmark 8.4' is a live Excel formula that reads the", False),
        ("yellow cells in 'Inputs'. Nothing is pasted from Python. Change an input and Excel recalculates.", False),
        ("Column C of 'Benchmark 8.4' is the spec table typed by hand. Column E is the Python engine's value,", False),
        ("written when the file was generated (python -m analyser.cashflow.benchmark, from backend/).", False),
        ("Columns F and G say MATCH when Excel and Python agree with the spec to 4 decimals.", False),
        ("", False),
        ("Conventions (same as the Python engine, see backend/analyser/cashflow/FORMULAS.md)", True),
        ("Units: Rs crore. Row = end of month. Month 0 = signing. Work of month m is billed at the end of month m.", False),
        ("Retention = r x bill until cumulative retention reaches Rmax x V.", False),
        ("Advance recovery = rho x bill until the advance received so far is repaid; any balance is recovered in", False),
        ("full in month floor(0.8 x T) = 19 and in the final bill (month T).", False),
        ("Net certified is paid after the payment lag (30 days = 1 month). Advances arrive in their months.", False),
        ("Delay D moves completion to month T + D/30; damages are deducted then; retention is released", False),
        ("'RelMonths' after completion. Cash out = cost share x gross work, paid the same month.", False),
        ("Values not given in spec §8.4 are [ASSUMPTION]s and are marked in 'Inputs'.", False),
    ]
    for i, (text, bold) in enumerate(lines, start=1):
        c = ws.cell(row=i, column=1, value=text)
        if bold:
            c.font = BOLD
    ws["A3"].fill = WARN_FILL
    ws.column_dimensions["A"].width = 110


def write_workbook(path: Path | str = WORKBOOK_PATH) -> Path:
    """Write the benchmark workbook with real Excel formulas. Returns the path."""
    path = Path(path)
    py = engine_figures()
    wb = Workbook()
    _readme(wb.active)
    wb.active.title = "Read me"

    # Inputs sheet with a defined name per input
    wsi = wb.create_sheet("Inputs")
    for col, h in enumerate(["Name", "Input", "Value", "Unit / note"], start=1):
        wsi.cell(row=1, column=col, value=h).font = BOLD
    for i, (name, label, value, unit) in enumerate(INPUTS, start=2):
        wsi.cell(row=i, column=1, value=name)
        wsi.cell(row=i, column=2, value=label)
        c = wsi.cell(row=i, column=3, value=value)
        c.fill = INPUT_FILL
        wsi.cell(row=i, column=4, value=unit)
        wb.defined_names[name] = DefinedName(name, attr_text=f"Inputs!$C${i}")
    for col, w in zip("ABCD", (14, 38, 10, 52)):
        wsi.column_dimensions[col].width = w

    # Monthly sheet
    wsm = wb.create_sheet("Monthly")
    for col, (h, _) in enumerate(MONTH_COLS, start=1):
        c = wsm.cell(row=1, column=col, value=h)
        c.font, c.fill = BOLD, HEAD_FILL
        c.alignment = Alignment(wrap_text=True)
        wsm.column_dimensions[c.column_letter].width = 13
    for m in range(0, LAST_MONTH + 1):
        r = FIRST_ROW + m
        wsm.cell(row=r, column=1, value=m)
        for col, (_, tpl) in enumerate(MONTH_COLS[1:], start=2):
            if m == 0 and "{p}" in tpl:
                # month 0 has no previous row: cumulative columns start from this row's own value
                tpl = (tpl.replace("D{p}", "0").replace("F{p}", "0").replace("H{p}", "0")
                          .replace("O{p}", "0").replace("P{p}", "0"))
            wsm.cell(row=r, column=col, value=tpl.format(r=r, p=r - 1, ld=LD_ROW))
            wsm.cell(row=r, column=col).number_format = "0.0000"
        wsm.cell(row=r, column=1).number_format = "0"
    wsm.freeze_panes = "B2"

    # Benchmark sheet
    wsb = wb.create_sheet("Benchmark 8.4")
    heads = ["Item", "Hand calculation", "Spec §8.4 result (typed)", "Excel formula result",
             "Python engine", "Excel = spec?", "Python = spec?"]
    for col, h in enumerate(heads, start=1):
        c = wsb.cell(row=1, column=col, value=h)
        c.font, c.fill = BOLD, HEAD_FILL
    last = FIRST_ROW + LAST_MONTH
    excel = {
        "gross_month": "=Monthly!B3",
        "retention_month": "=Monthly!C3",
        "retention_cap_month": f"=MIN(Monthly!S2:S{last})",
        "recovery_month": "=Monthly!G3",
        "advance_repaid_month": f"=MIN(Monthly!T2:T{last})",
        "deadline_month": "=RecDeadline*Months",
        "net_month1": "=Monthly!J3",
        "ld_60": "=MIN(LDRate*DelayDays,LDCap)*ContractValue",
        "ld_days_to_cap": "=LDCap/LDRate",
        "ld_at_cap": "=LDCap*ContractValue",
    }
    match = '=IF(ABS(ROUND({x}{r},4)-C{r})<0.00005,"MATCH","CHECK")'
    for i, (key, item, hand, val) in enumerate(SPEC_TABLE, start=2):
        wsb.cell(row=i, column=1, value=item)
        wsb.cell(row=i, column=2, value=hand)
        wsb.cell(row=i, column=3, value=val)
        wsb.cell(row=i, column=4, value=excel[key])
        wsb.cell(row=i, column=5, value=round(py[key], 10))
        wsb.cell(row=i, column=6, value=match.format(x="D", r=i))
        wsb.cell(row=i, column=7, value=match.format(x="E", r=i))

    r0 = len(SPEC_TABLE) + 4
    wsb.cell(row=r0, column=1, value="Beyond the spec table (Excel vs Python only; no spec value)").font = BOLD
    extra = [
        ("peak_funding", "Peak funding need = largest positive gap, 0 if none (Rs crore)",
         f"=MAX(0,MAX(Monthly!Q2:Q{last}))"),
        ("peak_month", "Month of peak funding (0 if never short of cash)",
         f"=IF(MAX(Monthly!Q2:Q{last})<=0,0,"
         f"INDEX(Monthly!A2:A{last},MATCH(MAX(Monthly!Q2:Q{last}),Monthly!Q2:Q{last},0)))"),
        ("overdraft_interest", "Overdraft interest = sum of positive gaps x rate / 12 (Rs crore)",
         f'=SUMIF(Monthly!Q2:Q{last},">0")*ODRate/12'),
        ("total_cash_in", "Total cash in = V - delay damages (Rs crore)", f"=Monthly!O{last}"),
        ("total_cash_out", "Total cash out = c x V (Rs crore)", f"=Monthly!P{last}"),
    ]
    wsb.cell(row=r0 + 1, column=1, value="Item").font = BOLD
    wsb.cell(row=r0 + 1, column=4, value="Excel formula result").font = BOLD
    wsb.cell(row=r0 + 1, column=5, value="Python engine").font = BOLD
    wsb.cell(row=r0 + 1, column=6, value="Excel = Python?").font = BOLD
    for j, (key, item, f) in enumerate(extra, start=r0 + 2):
        wsb.cell(row=j, column=1, value=item)
        wsb.cell(row=j, column=4, value=f)
        wsb.cell(row=j, column=5, value=round(py[key], 10))
        wsb.cell(row=j, column=6, value=f'=IF(ABS(D{j}-E{j})<0.00005,"MATCH","CHECK")')
    for col, w in zip("ABCDEFG", (62, 24, 22, 20, 16, 14, 14)):
        wsb.column_dimensions[col].width = w
    for row in wsb.iter_rows(min_row=2, min_col=3, max_col=5):
        for c in row:
            c.number_format = "0.0000"

    # Python monthly values side by side with the Excel gap (values, then a live difference formula)
    wsp = wb.create_sheet("Python monthly")
    heads = ["Month", "Python cash in", "Python cash out", "Python gap", "Excel gap - Python gap"]
    for col, h in enumerate(heads, start=1):
        c = wsp.cell(row=1, column=col, value=h)
        c.font, c.fill = BOLD, HEAD_FILL
        wsp.column_dimensions[c.column_letter].width = 22
    res = run_cashflow(benchmark_inputs(), sensitivities=False)
    for row in res.months:
        r = FIRST_ROW + row.month
        wsp.cell(row=r, column=1, value=row.month)
        wsp.cell(row=r, column=2, value=round(row.cash_in / CRORE, 10))
        wsp.cell(row=r, column=3, value=round(row.cash_out / CRORE, 10))
        wsp.cell(row=r, column=4, value=round(row.gap / CRORE, 10))
        wsp.cell(row=r, column=5, value=f"=Monthly!Q{r}-D{r}")
    j = r0 + 2 + len(extra)
    wsb.cell(row=j, column=1, value="Largest monthly difference, Excel gap vs Python gap (Rs crore)")
    wsb.cell(row=j, column=4, value=f"=MAX(ABS(MAX('Python monthly'!E2:E{last})),ABS(MIN('Python monthly'!E2:E{last})))")
    wsb.cell(row=j, column=5, value=0)
    wsb.cell(row=j, column=6, value=f'=IF(D{j}<0.00005,"MATCH","CHECK")')

    wb.move_sheet("Benchmark 8.4", offset=-1)   # order: Read me, Inputs, Benchmark 8.4, Monthly, Python monthly
    wb.move_sheet("Python monthly", offset=1)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


if __name__ == "__main__":
    out = write_workbook()
    print(f"Wrote {out}")
    fig = engine_figures()
    for key, item, hand, val in SPEC_TABLE:
        print(f"{item:60s} spec {val:>9}  engine {fig[key]:.4f}")
